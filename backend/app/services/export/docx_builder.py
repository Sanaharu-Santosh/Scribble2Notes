"""Scene -> Word document.

The goal is a document that is *editable in Word*, not a picture of one. So
headings become real headings, tables become real Word tables with addressable
cells, and positions are thrown away in favour of reading order — which is what
makes the result behave like something you typed rather than something you
scanned.

Known limitation, worth stating plainly: a mark applies to the whole paragraph.
The canvas knows an underline runs from x=244 to x=634, but not which characters
those pixels sit under — mapping the span back onto the text would need per-word
coordinates that survive the user retyping the paragraph. Underlining the whole
paragraph is wrong in a way the user can see and fix in two clicks; guessing a
character range is wrong in a way they would have to hunt for.
"""

from __future__ import annotations

import colorsys
import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

from app.schemas.export import ExportDocument, ExportItem, ItemKind

# Word offers only a fixed highlight palette, so the page's marker colour is
# matched to the nearest of these. Matched by HUE, not by RGB distance: real
# highlighter ink is pale, and grey sits in the middle of the RGB cube, so
# straight Euclidean distance sends most actual highlights to grey. A pale green
# (124, 255, 124) is closer to (128, 128, 128) than to (0, 255, 0) by that
# measure, which is arithmetically true and visibly wrong.
HIGHLIGHT_HUES: list[tuple[float, WD_COLOR_INDEX]] = [
    (0.0, WD_COLOR_INDEX.RED),
    (60.0, WD_COLOR_INDEX.YELLOW),
    (120.0, WD_COLOR_INDEX.BRIGHT_GREEN),
    (180.0, WD_COLOR_INDEX.TURQUOISE),
    (240.0, WD_COLOR_INDEX.BLUE),
    (300.0, WD_COLOR_INDEX.PINK),
]

# Below this saturation there is no hue worth matching — it really is grey.
MIN_SATURATION = 0.2


def _parse_hex(color: str | None) -> tuple[int, int, int] | None:
    if not color:
        return None
    value = color.strip().lstrip("#")
    if len(value) == 3:
        value = "".join(char * 2 for char in value)
    if len(value) != 6:
        return None
    try:
        return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
    except ValueError:
        return None


def nearest_highlight(color: str | None) -> WD_COLOR_INDEX:
    rgb = _parse_hex(color)
    if rgb is None:
        return WD_COLOR_INDEX.YELLOW  # the colour of most highlighters

    red, green, blue = (channel / 255 for channel in rgb)
    hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
    if saturation < MIN_SATURATION or lightness < 0.1:
        return WD_COLOR_INDEX.GRAY_25

    degrees = hue * 360

    def gap(candidate: float) -> float:
        """Hue is a circle: 350° and 10° are twenty degrees apart, not 340."""
        difference = abs(candidate - degrees) % 360
        return min(difference, 360 - difference)

    return min(HIGHLIGHT_HUES, key=lambda entry: gap(entry[0]))[1]


def _add_text(document: Document, item: ExportItem) -> None:
    text = (item.text or "").strip()
    if not text:
        return

    kinds = {mark.kind for mark in item.marks}

    if item.kind is ItemKind.HEADING:
        paragraph = document.add_heading(level=1)
    else:
        paragraph = document.add_paragraph()

    run = paragraph.add_run(text)
    if "underline" in kinds:
        run.underline = True
    if "highlight" in kinds:
        highlight = next(mark for mark in item.marks if mark.kind == "highlight")
        run.font.highlight_color = nearest_highlight(highlight.color)
    if item.kind is not ItemKind.HEADING and item.font_size:
        # Canvas font sizes are page pixels; Word wants points. Scanned pages
        # here are ~150dpi, so px * 72/150 lands close to the original size.
        run.font.size = Pt(max(round(item.font_size * 0.48), 8))


def _add_table(document: Document, item: ExportItem) -> None:
    if not item.table or not item.table.cells:
        return

    table = document.add_table(rows=item.table.rows, cols=item.table.cols)
    table.style = "Table Grid"

    for cell in item.table.cells:
        if 0 <= cell.row < item.table.rows and 0 <= cell.col < item.table.cols:
            table.cell(cell.row, cell.col).text = cell.text

    # A header row is a reasonable default for a ruled table and costs nothing
    # if wrong — the user can clear it in Word.
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def build_docx(document_model: ExportDocument) -> bytes:
    document = Document()

    if document_model.title:
        document.add_heading(document_model.title, level=0)

    for item in document_model.in_reading_order():
        if item.kind is ItemKind.TABLE:
            _add_table(document, item)
        elif item.kind in {ItemKind.HEADING, ItemKind.PARAGRAPH} or (
            # A boxed region only earns a paragraph if something is written in it.
            item.kind is ItemKind.FIGURE and item.text
        ):
            _add_text(document, item)
        # Lines and arrows are spatial marks with no place in a flowed
        # document; they survive in the PDF export instead.

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
