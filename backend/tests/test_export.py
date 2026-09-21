"""Export tests that open the files, rather than checking a byte count.

The two exports make different promises, and each is asserted here:

* the DOCX is *editable* — real headings, real Word tables with addressable
  cells, runs that carry the marks;
* the PDF is *searchable* — its text extracts as text, which is the whole
  difference between this and screenshotting the canvas.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfReader

from app.schemas.export import ExportDocument, ExportItem, ExportMark, ExportTable, ItemKind
from app.services.export.docx_builder import build_docx, nearest_highlight
from app.services.export.pdf_builder import PT_PER_PX, build_pdf, render_html, safe_color

PAGE_WIDTH, PAGE_HEIGHT = 1240, 1754


def _document(**overrides) -> ExportDocument:
    items = [
        ExportItem(
            id="s2n:heading:b0",
            kind=ItemKind.HEADING,
            x=100,
            y=110,
            width=800,
            height=60,
            reading_order=0,
            text="Unit 4 - Layout analysis",
        ),
        ExportItem(
            id="s2n:paragraph:b1",
            kind=ItemKind.PARAGRAPH,
            x=100,
            y=260,
            width=900,
            height=60,
            reading_order=1,
            text="Underlined and highlighted body text.",
            font_size=34,
            marks=[
                ExportMark(kind="underline", color="#1a1a1a", x=100, y=315, width=390, height=5),
                ExportMark(kind="highlight", color="#ffe25c", x=280, y=326, width=365, height=48),
            ],
        ),
        ExportItem(
            id="s2n:table:b2",
            kind=ItemKind.TABLE,
            x=100,
            y=450,
            width=1040,
            height=380,
            reading_order=2,
            table=ExportTable(
                rows=2,
                cols=2,
                cells=[
                    {"row": 0, "col": 0, "text": "Region"},
                    {"row": 0, "col": 1, "text": "Becomes"},
                    {"row": 1, "col": 0, "text": "Table"},
                    {"row": 1, "col": 1, "text": "real grid"},
                ],
            ),
        ),
        ExportItem(
            id="user-arrow-1",
            kind=ItemKind.ARROW,
            x=200,
            y=900,
            width=200,
            height=50,
            reading_order=3,
            color="#d63a2e",
            points=[(0, 0), (200, 50)],
        ),
    ]
    return ExportDocument(
        page_width=PAGE_WIDTH,
        page_height=PAGE_HEIGHT,
        title="My notes",
        items=items,
        **overrides,
    )


# --------------------------------------------------------------------------
# DOCX — is it editable?
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def docx_document() -> Document:
    return Document(io.BytesIO(build_docx(_document())))


def test_docx_keeps_headings_as_headings(docx_document):
    """Not bold text pretending to be a heading — a real style, so Word's
    navigation pane and table of contents work."""
    headings = [p for p in docx_document.paragraphs if p.style.name.startswith("Heading")]
    assert any("Layout analysis" in p.text for p in headings)


def test_docx_table_is_a_real_table(docx_document):
    assert len(docx_document.tables) == 1
    table = docx_document.tables[0]

    assert (len(table.rows), len(table.columns)) == (2, 2)
    assert table.cell(0, 0).text == "Region"
    assert table.cell(1, 1).text == "real grid"


def test_docx_carries_the_marks_onto_the_run(docx_document):
    body = next(p for p in docx_document.paragraphs if "Underlined" in p.text)
    run = body.runs[0]

    assert run.underline is True
    assert run.font.highlight_color is not None


def test_docx_follows_reading_order(docx_document):
    texts = [p.text for p in docx_document.paragraphs if p.text.strip()]
    assert texts.index("Unit 4 - Layout analysis") < texts.index(
        "Underlined and highlighted body text."
    )


def test_docx_drops_strokes_that_have_no_place_in_a_flowed_document(docx_document):
    """An arrow is a spatial mark. It belongs in the PDF, not as a stray
    paragraph in a Word document."""
    assert all("arrow" not in p.text.lower() for p in docx_document.paragraphs)


@pytest.mark.parametrize(
    ("color", "expected"),
    [
        ("#ffe25c", "YELLOW"),
        ("#7cff7c", "BRIGHT_GREEN"),
        ("#9ad8ff", "TURQUOISE"),
        ("#ff9ad8", "PINK"),
        ("#cccccc", "GRAY_25"),
        (None, "YELLOW"),
        ("nonsense", "YELLOW"),
    ],
)
def test_highlight_colour_maps_to_words_fixed_palette(color, expected):
    """Word only has named highlight colours, so the page's marker colour is
    matched to the nearest rather than reproduced.

    Matched by hue, because highlighter ink is pale and grey sits in the middle
    of the RGB cube: by plain RGB distance a pale green is nearer grey than
    green, so every real highlight would export grey.
    """
    assert nearest_highlight(color).name == expected


# --------------------------------------------------------------------------
# PDF — is it searchable, and is it the right size?
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def pdf_bytes() -> bytes:
    return build_pdf(_document())


def test_pdf_text_is_real_text(pdf_bytes):
    """The promise that separates this from exporting a PNG of the canvas."""
    text = PdfReader(io.BytesIO(pdf_bytes)).pages[0].extract_text()

    assert "Layout analysis" in text
    assert "Underlined and highlighted body text." in text
    assert "real grid" in text  # table cells too, not just paragraphs


def test_pdf_page_matches_the_scanned_page(pdf_bytes):
    """A 150dpi A4 scan should come out as an A4 PDF, not a 13-inch monster."""
    box = PdfReader(io.BytesIO(pdf_bytes)).pages[0].mediabox

    assert float(box.width) == pytest.approx(PAGE_WIDTH * PT_PER_PX, abs=2)
    assert float(box.height) == pytest.approx(PAGE_HEIGHT * PT_PER_PX, abs=2)
    assert float(box.width) == pytest.approx(595, abs=3)  # A4 width in points


def test_pdf_places_items_where_the_canvas_had_them(pdf_bytes):
    html = render_html(_document())
    # The heading sits at y=110px, which is 110 * 72/150 = 52.8pt.
    assert "top:52.80pt" in html


def test_strokes_survive_into_the_pdf():
    html = render_html(_document())
    assert "<svg" in html and "polyline" in html


# --------------------------------------------------------------------------
# The document is user content, and it ends up inside HTML
# --------------------------------------------------------------------------


def test_text_is_escaped_not_interpreted():
    document = _document()
    document.items[0].text = "<script>alert(1)</script> & co"

    html = render_html(document)

    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "&amp; co" in html


@pytest.mark.parametrize(
    "hostile",
    ["red; position:fixed", "#fff\" onload=\"x", "url(javascript:alert(1))", "expression(1)"],
)
def test_colours_that_are_not_colours_are_dropped(hostile):
    """Colours come from the canvas, so from the user, and land in a style
    attribute. Nothing that isn't plainly a colour gets written."""
    assert safe_color(hostile) == "#1a2233"


def test_ordinary_colours_pass_through():
    assert safe_color("#ffe25c") == "#ffe25c"
    assert safe_color("#fff") == "#fff"
    assert safe_color("transparent", "transparent") == "transparent"


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------


def test_docx_endpoint_returns_a_word_attachment(client):
    response = client.post("/api/export/docx", json=_document().model_dump(mode="json"))

    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert 'filename="My-notes.docx"' in response.headers["content-disposition"]
    assert Document(io.BytesIO(response.content)).tables


def test_pdf_endpoint_returns_a_pdf_attachment(client):
    response = client.post("/api/export/pdf", json=_document().model_dump(mode="json"))

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_filenames_from_page_titles_cannot_escape(client):
    document = _document()
    document.title = "../../etc/passwd"

    response = client.post("/api/export/pdf", json=document.model_dump(mode="json"))

    disposition = response.headers["content-disposition"]
    assert ".." not in disposition
    assert "/" not in disposition.split("filename=")[1]


def test_an_empty_page_still_exports(client):
    empty = ExportDocument(page_width=PAGE_WIDTH, page_height=PAGE_HEIGHT, items=[])

    response = client.post("/api/export/docx", json=empty.model_dump(mode="json"))

    assert response.status_code == 200
    assert 'filename="notes.docx"' in response.headers["content-disposition"]
